# -*- coding: utf-8 -*-
"""D(#10/#11): materialize 範囲限定と skip_hidden 時の read_only 維持。"""
from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import openpyxl

from svc import svc_data_agg_extract as ex
from svc.data_agg_row_visibility import get_hidden_excel_rows


def test_ensure_hidden_wb_does_not_evict_readonly(tmp_path: Path) -> None:
    p = tmp_path / "h.xlsx"
    wb0 = openpyxl.Workbook()
    wb0.active.title = "S"
    wb0.active["A1"] = "v"
    wb0.active.row_dimensions[2].hidden = True
    wb0.save(p)
    wb0.close()

    with ex.xlsx_workbook_scope():
        ro = ex._xlsx_workbook_from_cache(p)
        assert ro is not None
        assert bool(getattr(ro, "read_only", False)) is True
        got = ex.ensure_xlsx_workbook_for_hidden_rows(p)
        assert got is None  # read_only を潰さない
        still = ex._xlsx_workbook_from_cache(p)
        assert still is ro
        assert bool(getattr(still, "read_only", False)) is True
        # 非表示判定は短寿命経路で取得できる
        hidden = get_hidden_excel_rows(p, "S")
        assert 1 in hidden


def test_ensure_returns_existing_non_readonly(tmp_path: Path) -> None:
    p = tmp_path / "nr.xlsx"
    wb0 = openpyxl.Workbook()
    wb0.active.title = "S"
    wb0.active["A1"] = "v"
    wb0.save(p)
    wb0.close()

    with ex.xlsx_workbook_scope():
        frame = ex._xlsx_workbook_cache_top()
        assert frame is not None
        full = openpyxl.load_workbook(p, read_only=False, data_only=True)
        frame.setdefault("wbs", {})[ex._xlsx_cache_path_key(p)] = full
        got = ex.ensure_xlsx_workbook_for_hidden_rows(p)
        assert got is full
        assert bool(getattr(got, "read_only", False)) is False


def test_materialize_respects_used_bounds() -> None:
    class _Ws:
        max_row = 2
        max_column = 2

        def iter_rows(self, **kwargs):
            assert kwargs.get("max_row") == 2
            assert kwargs.get("max_col") == 2
            yield (SimpleNamespace(value="a"), SimpleNamespace(value="b"))
            yield (SimpleNamespace(value="c"), SimpleNamespace(value="d"))

    mat = ex._materialize_readonly_sheet_matrix(_Ws())
    assert len(mat) == 2
    assert len(mat[0]) == 2


def test_materialize_falls_back_when_bounds_unknown() -> None:
    class _Ws:
        max_row = 0
        max_column = 0

        def iter_rows(self, values_only=False, **kwargs):
            assert values_only is False
            assert "max_row" not in kwargs
            yield (SimpleNamespace(value="x"),)

    mat = ex._materialize_readonly_sheet_matrix(_Ws())
    assert len(mat) == 1


def test_multi_sheet_hidden_loads_full_workbook_once(
    tmp_path: Path, monkeypatch
) -> None:
    """スコープ内で複数シートの skip_hidden でも full open はファイルあたり1回。"""
    p = tmp_path / "multi.xlsx"
    wb0 = openpyxl.Workbook()
    wb0.active.title = "S1"
    wb0.active["A1"] = "a"
    wb0.active.row_dimensions[2].hidden = True
    for name in ("S2", "S3"):
        ws = wb0.create_sheet(name)
        ws["A1"] = "b"
        ws.row_dimensions[3].hidden = True
    wb0.save(p)
    wb0.close()

    full_opens = {"n": 0}
    real_load = openpyxl.load_workbook

    def _counting_load(*args, **kwargs):
        if not bool(kwargs.get("read_only", False)):
            full_opens["n"] += 1
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _counting_load)

    with ex.xlsx_workbook_scope():
        ro = ex._xlsx_workbook_from_cache(p)
        assert ro is not None
        assert bool(getattr(ro, "read_only", False)) is True
        h1 = get_hidden_excel_rows(p, "S1")
        h2 = get_hidden_excel_rows(p, "S2")
        h3 = get_hidden_excel_rows(p, "S3")
        assert 1 in h1
        assert 2 in h2
        assert 2 in h3
        still = ex._xlsx_workbook_from_cache(p)
        assert still is ro

    assert full_opens["n"] == 1
