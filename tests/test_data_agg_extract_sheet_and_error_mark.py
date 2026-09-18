# -*- coding: utf-8 -*-
"""A(#2/#3): シート欠落エラー化と抽出失敗マーカー。"""
from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from svc.data_agg_sheet_resolve import (
    EXTRACT_READ_ERROR_MARK,
    DataAggSheetMissingError,
    is_extract_read_error,
)
from svc.svc_data_agg_extract import (
    _get_excel_cell,
    _resolve_readonly_worksheet,
    _xlsx_cell_value_open_workbook_rc,
)


def _xlsx_two_sheets(tmp_path: Path) -> Path:
    p = tmp_path / "two.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "on_sheet1"
    ws2 = wb.create_sheet("Sheet2")
    ws2["A1"] = "on_sheet2"
    wb.save(p)
    wb.close()
    return p


def test_resolve_missing_sheet_raises() -> None:
    wb = SimpleNamespace(sheetnames=["A", "B"], active=SimpleNamespace(title="A"))
    with pytest.raises(DataAggSheetMissingError) as ei:
        _resolve_readonly_worksheet(wb, "NoSuch")
    assert "NoSuch" in str(ei.value)


def test_resolve_empty_sheet_uses_active() -> None:
    active = SimpleNamespace(title="Left")
    wb = SimpleNamespace(sheetnames=["Left"], active=active)
    ws, name = _resolve_readonly_worksheet(wb, None)
    assert ws is active
    assert name == "Left"
    ws2, name2 = _resolve_readonly_worksheet(wb, "  ")
    assert ws2 is active
    assert name2 == "Left"


def test_get_excel_cell_missing_sheet_raises(tmp_path: Path) -> None:
    p = _xlsx_two_sheets(tmp_path)
    with pytest.raises(DataAggSheetMissingError):
        _get_excel_cell(p, "MissingSheet", "A1")


def test_get_excel_cell_existing_sheet_ok(tmp_path: Path) -> None:
    p = _xlsx_two_sheets(tmp_path)
    assert _get_excel_cell(p, "Sheet2", "A1") == "on_sheet2"
    assert _get_excel_cell(p, "Sheet1", "A1") == "on_sheet1"


def test_extract_read_error_mark_is_err_extract() -> None:
    assert EXTRACT_READ_ERROR_MARK == "#ERR_EXTRACT"
    assert is_extract_read_error("#ERR_EXTRACT")
    assert not is_extract_read_error("（抽出失敗）")
    assert not is_extract_read_error(None)
    assert not is_extract_read_error("")


def test_extract_item_bundle_missing_sheet_skips_empty(tmp_path: Path) -> None:
    """指定シート無しは一括入口で空スキップ。印は入れず、例外でも止めない。"""
    from svc.svc_data_agg_extract import extract_item_bundle

    p = _xlsx_two_sheets(tmp_path)
    item = {
        "sources": [
            {
                "type": "cell",
                "sheet_name": "NoSuch",
                "cell_ref": "A1",
                "ui_scenario_source_v1": {"sheet_rule": "完全一致"},
            }
        ]
    }
    out = extract_item_bundle(p, item)
    assert out.get("primary_values") == []
    assert "#ERR_EXTRACT" not in (out.get("primary_values") or [])


def test_cell_read_exception_returns_extract_mark() -> None:
    class _BoomWs:
        def __getitem__(self, _ref: str) -> None:
            raise RuntimeError("boom")

    class _BoomWb:
        sheetnames = ["S"]

        def __getitem__(self, name: str) -> _BoomWs:
            assert name == "S"
            return _BoomWs()

    out = _xlsx_cell_value_open_workbook_rc(_BoomWb(), "S", 0, 0)
    assert out == EXTRACT_READ_ERROR_MARK
    assert is_extract_read_error(out)
    assert not is_extract_read_error(None)
    assert not is_extract_read_error("")
