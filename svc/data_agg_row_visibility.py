# -*- coding: utf-8 -*-
"""Excel 行の非表示（手動・オートフィルタ）判定。抽出の skip_hidden_rows 用。"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Callable, Optional

from svc.data_agg_path_norm import normalize_source_path_literal

logger = logging.getLogger(__name__)

_CACHE_KEY = "hidden_row_sets"  # frame[path][sheet_key] -> set[int] (0-based) | frozenset


def source_wants_skip_hidden_rows(src: dict[str, Any] | None) -> bool:
    """ソースで非表示・フィルタ行を走査から除外するか（既定 OFF）。"""
    if not isinstance(src, dict):
        return False
    v = src.get("skip_hidden_rows")
    if v is True or v == 1:
        return True
    if isinstance(v, str) and v.strip().lower() in ("1", "true", "yes", "on"):
        return True
    return False


def _sheet_cache_key(sheet_name: Optional[str]) -> str:
    return str(sheet_name or "").strip() or "__LEFTMOST__"


def _frame_hidden_cache() -> dict[str, Any] | None:
    try:
        from svc.svc_data_agg_extract import _xlsx_workbook_cache_top
    except Exception:
        return None
    frame = _xlsx_workbook_cache_top()
    if frame is None:
        return None
    return frame.setdefault(_CACHE_KEY, {})


def _hidden_set_from_worksheet(ws: Any) -> set[int]:
    """1 シートの row_dimensions.hidden を 0 始まり集合にする。"""
    hidden: set[int] = set()
    if ws is None:
        return hidden
    dims = getattr(ws, "row_dimensions", None)
    if dims is None:
        return hidden
    for idx, dim in dims.items():
        try:
            r1 = int(idx)
        except (TypeError, ValueError):
            continue
        if bool(getattr(dim, "hidden", False)) and r1 >= 1:
            hidden.add(r1 - 1)
    return hidden


def _store_all_sheets_hidden(wb: Any, by_sheet: dict[str, Any]) -> None:
    """
    1 回開いた（またはキャッシュ済みの）非 read_only ブックから、
    全シート分の非表示行を by_sheet に格納する（追加の load_workbook を避ける）。
    """
    if wb is None or by_sheet is None:
        return
    names = [str(x) for x in (getattr(wb, "sheetnames", None) or []) if str(x).strip()]
    for name in names:
        try:
            ws = wb[name]
        except Exception:
            by_sheet[_sheet_cache_key(name)] = frozenset()
            continue
        by_sheet[_sheet_cache_key(name)] = frozenset(_hidden_set_from_worksheet(ws))
    active = getattr(wb, "active", None)
    active_title = str(getattr(active, "title", "") or "").strip() if active is not None else ""
    if active_title and _sheet_cache_key(active_title) in by_sheet:
        by_sheet["__LEFTMOST__"] = by_sheet[_sheet_cache_key(active_title)]
    elif active is not None:
        left = frozenset(_hidden_set_from_worksheet(active))
        by_sheet["__LEFTMOST__"] = left
        if active_title:
            by_sheet.setdefault(_sheet_cache_key(active_title), left)


def _hidden_for_request(
    wb: Any,
    sheet_name: Optional[str],
    *,
    by_sheet: dict[str, Any] | None,
) -> set[int]:
    """wb から要求シートの hidden を返す。by_sheet があれば全シート分を一度に埋める。"""
    sk = _sheet_cache_key(sheet_name)
    if by_sheet is not None:
        if sk not in by_sheet:
            _store_all_sheets_hidden(wb, by_sheet)
        if sk not in by_sheet:
            sn = str(sheet_name or "").strip()
            if sn:
                logger.warning(
                    "[DATA_AGG_HIDDEN] シート「%s」なし — 非表示判定をスキップ",
                    sn,
                )
            by_sheet[sk] = frozenset()
        return set(by_sheet[sk])

    sn = str(sheet_name or "").strip()
    names = [str(x) for x in (getattr(wb, "sheetnames", None) or [])]
    try:
        if sn:
            if sn not in names:
                logger.warning(
                    "[DATA_AGG_HIDDEN] シート「%s」なし — 非表示判定をスキップ",
                    sn,
                )
                return set()
            ws = wb[sn]
        else:
            ws = getattr(wb, "active", None)
    except Exception as e:
        logger.debug("[DATA_AGG_HIDDEN] wb sheet 解決失敗: %s", e)
        return set()
    return _hidden_set_from_worksheet(ws)


def _load_hidden_rows_xlsx(path: Path, sheet_name: Optional[str]) -> set[int]:
    """
    openpyxl 通常ロードで非表示行を収集（0 始まり行）。

    xlsx_workbook_scope 内では 1 回の open で全シート分を hidden キャッシュへ格納し、
    複数シートの skip_hidden でも full オープンがシート数分に増えないようにする。
    read_only 抽出キャッシュは触らない。
    """
    import time

    from core import core_env

    t0 = time.perf_counter()
    by_sheet_root = _frame_hidden_cache()
    pk = _path_cache_key(path)
    by_sheet = by_sheet_root.setdefault(pk, {}) if by_sheet_root is not None else None

    try:
        import openpyxl
    except ImportError:
        return set()
    wb = None
    try:
        wb = openpyxl.load_workbook(
            path, read_only=False, data_only=False, keep_links=False
        )
    except Exception as e:
        logger.debug("[DATA_AGG_HIDDEN] xlsx open 失敗 %s: %s", path, e)
        return set()
    try:
        return _hidden_for_request(wb, sheet_name, by_sheet=by_sheet)
    except Exception as e:
        logger.debug("[DATA_AGG_HIDDEN] xlsx dims 失敗 %s: %s", path, e)
        return set()
    finally:
        try:
            if wb is not None:
                wb.close()
        except Exception:
            pass
        if core_env.data_agg_io_profile_enabled():
            try:
                from svc import data_agg_io_profile_temp as iop

                iop.record_hidden_open(path, time.perf_counter() - t0)
            except Exception:
                pass


def _load_hidden_rows_xls(path: Path, sheet_name: Optional[str]) -> set[int]:
    """xlrd formatting_info=True で rowinfo_map.hidden を収集（0 始まり行）。"""
    try:
        import xlrd
    except ImportError:
        return set()
    hidden: set[int] = set()
    try:
        book = xlrd.open_workbook(
            str(path), on_demand=False, formatting_info=True
        )
    except Exception as e:
        logger.debug("[DATA_AGG_HIDDEN] xls open 失敗 %s: %s", path, e)
        return set()
    try:
        names = list(book.sheet_names() or [])
        sn = str(sheet_name or "").strip()
        if sn and sn in names:
            sheet = book.sheet_by_name(sn)
        else:
            sheet = book.sheet_by_index(0)
        info_map = getattr(sheet, "rowinfo_map", None) or {}
        for r, info in info_map.items():
            try:
                ri = int(r)
            except (TypeError, ValueError):
                continue
            if bool(getattr(info, "hidden", False)):
                hidden.add(ri)
    except Exception as e:
        logger.debug("[DATA_AGG_HIDDEN] xls rowinfo 失敗 %s: %s", path, e)
    finally:
        try:
            release = getattr(book, "release_resources", None)
            if callable(release):
                release()
        except Exception:
            pass
    return hidden


def _path_cache_key(path: Path | str) -> str:
    return normalize_source_path_literal(path)


def _non_readonly_workbook_from_cache(path: Path) -> Any | None:
    """スコープ内の非 read_only ブック。read_only のみ／未登録なら None（潰さない）。"""
    try:
        from svc.svc_data_agg_extract import (  # noqa: WPS433
            ensure_xlsx_workbook_for_hidden_rows,
            is_openxml_excel_suffix,
        )
    except Exception:
        return None
    if not is_openxml_excel_suffix(path.suffix):
        return None
    return ensure_xlsx_workbook_for_hidden_rows(path)


def get_hidden_excel_rows(
    path: Path | str,
    sheet_name: Optional[str] = None,
) -> set[int]:
    """
    非表示（手動・フィルタで隠れた）行の 0 始まり集合。
    .csv や判定不能時は空集合。xlsx_workbook_scope 内ではキャッシュする。
    """
    p = Path(path)
    suf = p.suffix.lower()
    if suf not in (".xlsx", ".xlsm", ".xls"):
        return set()
    sk = _sheet_cache_key(sheet_name)
    pk = _path_cache_key(p)
    cache = _frame_hidden_cache()
    by_sheet: dict[str, Any] | None = None
    if cache is not None:
        by_sheet = cache.setdefault(pk, {})
        hit = by_sheet.get(sk)
        if hit is not None:
            return set(hit)
    if suf in (".xlsx", ".xlsm"):
        try:
            wb = _non_readonly_workbook_from_cache(p)
            if wb is not None:
                hidden = _hidden_for_request(wb, sheet_name, by_sheet=by_sheet)
            else:
                hidden = _load_hidden_rows_xlsx(p, sheet_name)
        except Exception:
            hidden = _load_hidden_rows_xlsx(p, sheet_name)
    else:
        hidden = _load_hidden_rows_xls(p, sheet_name)
    if by_sheet is not None:
        by_sheet.setdefault(sk, frozenset(hidden))
    return hidden


def make_row_hidden_predicate(
    path: Path | str,
    sheet_name: Optional[str],
    *,
    enabled: bool,
) -> Callable[[int], bool] | None:
    """
    Excel 行（0 始まり）が非表示なら True を返す述語。
    enabled=False または対象外形式なら None（スキップ無し）。
    """
    if not enabled:
        return None
    p = Path(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm", ".xls"):
        return None
    hidden = get_hidden_excel_rows(p, sheet_name)
    if not hidden:
        return None

    def _is_hidden(excel_row_0: int) -> bool:
        try:
            return int(excel_row_0) in hidden
        except (TypeError, ValueError):
            return False

    return _is_hidden
